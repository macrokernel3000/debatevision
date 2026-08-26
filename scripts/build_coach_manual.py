from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "辯語視界_八項活動教練手冊.docx"

FONT = "Heiti TC"
NAVY = RGBColor(26, 55, 77)
BLUE = RGBColor(39, 105, 139)
GOLD = RGBColor(187, 137, 48)
MUTED = RGBColor(92, 103, 112)
LIGHT_BLUE = "E8F1F5"
LIGHT_GOLD = "F8F0DD"
LIGHT_GRAY = "F3F5F6"
WHITE = RGBColor(255, 255, 255)


def set_run(run, size=11, bold=False, color=None, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=NAVY, size=9.6, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    set_run(p.add_run(text), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    grid = table._tbl.tblGrid
    for node in list(grid):
        grid.remove(node)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def add_text(doc, text="", size=11, bold=False, color=None, italic=False, after=6, align=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    set_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def add_labeled(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell_margins(cell, top=150, bottom=150, start=120, end=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    set_run(p.add_run(f"{label}｜"), size=10.5, bold=True, color=BLUE)
    set_run(p.add_run(text), size=10.5, color=NAVY)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.38)
        p.paragraph_format.first_line_indent = Inches(-0.19)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        set_run(p.add_run(item), size=10.5, color=NAVY)


def add_steps(doc, items):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.extend([tabs, ind])
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    for item in items:
        p = doc.add_paragraph()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num_id_node])
        p._p.get_or_add_pPr().append(num_pr)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        set_run(p.add_run(item), size=10.5, color=NAVY)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = level == 1
    p.paragraph_format.space_before = Pt(18 if level == 1 else 11)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 5)
    r = p.add_run(text)
    set_run(r, size={1: 18, 2: 14, 3: 11.5}[level], bold=True, color=NAVY if level == 1 else BLUE)
    return p


def add_example(doc, setup, answer, questions, takeaway):
    heading(doc, "完整示例", 2)
    add_labeled(doc, "抽到／設定", setup, LIGHT_GOLD)
    add_text(doc, "示範回答", size=11.5, bold=True, color=GOLD, after=3, keep=True)
    add_text(doc, answer, size=10.5, color=NAVY, after=6)
    add_text(doc, "教練追問", size=11.5, bold=True, color=BLUE, after=3, keep=True)
    add_bullets(doc, questions)
    add_labeled(doc, "收束重點", takeaway, LIGHT_GRAY)


ACTIVITIES = [
    {
        "title": "1｜異境求生",
        "tag": "限制思考・資源運用・團隊策略",
        "goal": ["辨認情境中的關鍵限制", "把普通資源轉化為具體用途", "比較不同方案的成本、風險與優先順序"],
        "time": "10–25 分鐘；個人、小組或全班競賽皆可。",
        "prepare": ["選擇求生版或冒險版。", "求生版決定抽道具、職業或兩者；冒險版設定隊伍數與每隊卡片類型。", "依年齡取消太陌生的卡牌；第一輪建議每人／每組 1–3 張。"],
        "steps": ["抽出同一個異境，請全班先說出最危險的三個限制。", "公布資源卡，給 30–60 秒討論。", "學生提出『做什麼、怎麼做、解決哪個限制』。", "其他組從太慢、太危險、資源不足或不適用提出一次質疑。", "原組補強後，全班依可行性投票。"],
        "coach": ["不要接受『這個很有用』，追問『在哪一步用？』", "允許小用途成立；重點是說得具體，不是卡牌本身強大。", "冒險版要學生說明隊員如何互補，不能只逐張介紹。"],
        "setup": "異境：無人島｜資源：垃圾袋、廚師",
        "answer": "我們先用垃圾袋收集雨水，避免直接飲用海水；廚師負責辨識食材、控制火候並分配食物。垃圾袋也能包住乾柴，讓點火材料不被午後雷雨打濕。第一天優先處理飲水與火，不急著蓋完整住所。",
        "questions": ["垃圾袋破掉怎麼辦？", "廚師真的會辨識野生植物嗎？", "如果三天都不下雨，方案要如何調整？"],
        "takeaway": "好答案會把資源放進行動順序，並主動承認限制；不是把每張卡說成萬能。",
    },
    {
        "title": "2｜現實召喚",
        "tag": "創意落地・制度理解・方案辯護",
        "goal": ["把超現實能力轉成現實可執行方案", "處理法律、成本、信任與能力限制", "提出有步驟、能驗證的計畫"],
        "time": "12–20 分鐘；適合個人短講或多角色方案競賽。",
        "prepare": ["抽 1 張現實任務，再抽 1–6 張召喚卡。", "第一輪建議 1 張角色，熟悉後再讓學生組隊。", "老師可鎖定任務，讓不同角色輪流挑戰同一問題。"],
        "steps": ["先確認任務的成功標準。", "學生說明角色的能力與最大現實障礙。", "提出第一步、合作對象、執行方法與驗證方式。", "同學從合法性、成本、可信度、替代方案追問。", "比較哪個方案最能把能力轉成制度內的成果。"],
        "coach": ["禁止只說『用魔法就完成』；能力只能是資源，不是答案。", "追問誰會允許他行動、誰付錢、誰相信結果。", "鼓勵學生設計小規模試辦，而非一步完成宏大目標。"],
        "setup": "任務：救一間學校｜召喚角色：讀心術能力者",
        "answer": "他不能直接讀全校師生的心，這會侵犯隱私。第一步由校方邀請自願參與的學生與老師，讓能力者只辨識匿名回饋中的情緒壓力，再交由輔導室整理三個最急迫問題。兩週後比較缺席率與求助人數，確認方案是否有效。",
        "questions": ["學生真的能自由拒絕嗎？", "讀到錯誤情緒時由誰負責？", "不用超能力，有沒有更便宜的替代方案？"],
        "takeaway": "創意不是跳過現實，而是讓特殊能力在法律、信任與驗證機制中產生價值。",
    },
    {
        "title": "3｜銷售密令",
        "tag": "需求洞察・敘事包裝・客群分析",
        "goal": ["區分產品功能與顧客需求", "針對不同對象調整賣點", "用證據回應價格與替代品質疑"],
        "time": "10–20 分鐘；可作暖身、簡報或小組競標。",
        "prepare": ["選供需版、故事版或目標版。", "供需版搭配需求；故事版搭配概念；目標版搭配動物、職業、名人或異族。", "商品一次抽 1–3 張最容易聚焦。"],
        "steps": ["說出顧客是誰、在什麼情境遇到什麼痛點。", "把商品功能翻譯成對顧客的利益。", "用一個具體使用畫面完成 45 秒提案。", "聽眾提出價格、替代品、可信度與必要性的質疑。", "銷售者修改一句核心賣點後再次提案。"],
        "coach": ["學生若只形容商品，追問『所以顧客得到什麼？』", "目標版不能依刻板印象亂猜，要引用卡牌身分或情境。", "故事不是越感人越好，必須回到購買理由。"],
        "setup": "目標版｜商品：保溫毯｜客群：消防員",
        "answer": "這不是賣給消防員自己蓋的毯子，而是讓第一線救援能立即降低傷者失溫風險。它重量輕、可大量放在消防車，不必等救護設備到場。每台車配置十件，比購買大型加熱設備更便宜，也能一次服務多人。",
        "questions": ["一般毛毯不能替代嗎？", "用完能否重複使用？", "哪種救援情境最需要它？"],
        "takeaway": "有說服力的銷售提案會同時回答：誰需要、何時需要、為何這件產品比替代品更適合。",
    },
    {
        "title": "4｜隱喻羅盤",
        "tag": "概念定義・關係建構・反例修正",
        "goal": ["找出兩個事物之間可說明的相似點", "重新定義抽象概念", "用例子與反例調整命題強度"],
        "time": "8–18 分鐘；適合語文暖身、寫作與論述訓練。",
        "prepare": ["初學者使用具體版：固定『人生就像』，選物品或動物後綴。", "進階使用抽象版：選前綴、關係、後綴牌庫。", "需要重複練習時，鎖定其中一個位置，只重抽其餘卡。"],
        "steps": ["先為兩端詞語各下簡短定義。", "找出一個共同結構或因果角度。", "用具體事件證明命題在某些情況成立。", "同學提出反例，指出哪個詞太模糊或關係太強。", "原說話者加上條件或改寫命題。"],
        "coach": ["不要問『真的一模一樣嗎』；比喻本來就是局部相似。", "抽象版先確定關係詞的強度：導致、需要、限制、促進並不相同。", "最好的收束常是把絕對句改成有條件的句子。"],
        "setup": "具體版：人生就像巧克力",
        "answer": "人生像一盒看不見內餡的巧克力：我們能選擇拿哪一顆，卻不能完全確定結果。有些經驗入口苦，但後來留下價值；重要的不是每次都選到喜歡的，而是遇到不喜歡時仍能決定下一步。",
        "questions": ["人生真的完全無法預測嗎？", "如果有人可以先看成分表，比喻還成立嗎？", "這個比喻忽略了人生哪一部分？"],
        "takeaway": "隱喻成立靠的是清楚的共同結構；接受反例後補上範圍，論述會比硬拗更有力量。",
    },
    {
        "title": "5｜誰更重要",
        "tag": "比較判準・攻防思考・立場轉換",
        "goal": ["先建立標準再做判斷", "區分短期與長期、個人與社會等不同尺度", "理解改變判準會改變結論"],
        "time": "8–15 分鐘；適合兩人對決、小組攻防或全班投票。",
        "prepare": ["紅角與藍角各選牌池，再各抽 1 張。", "初學者可用同類卡；進階可跨類別比較。", "提醒觀眾投票對象是『比較理由』，不是喜歡的卡。"],
        "steps": ["雙方各提出一個比較標準。", "各用 30 秒說明自己的卡在該標準下更重要。", "雙方各指出對方標準的一個限制。", "老師公布新情境或要求改變標準。", "觀眾說明採用哪個標準後投票。"],
        "coach": ["若學生說『本來就比較重要』，要求可共同檢驗的標準。", "允許雙方使用不同標準，但必須爭論哪個標準更適合題目。", "最後問一次：換到另一個情境，結果會不會翻盤？"],
        "setup": "紅角：醫生｜藍角：濾水器｜情境：地震後前三天",
        "answer": "我支持濾水器。前三天的標準應是『能同時降低多少人的立即生存風險』。醫生很重要，但沒有乾淨水源，傷患與健康居民都可能感染；濾水器能先減少新增病患，讓有限醫療人力留給重傷者。",
        "questions": ["沒有醫生，重傷者怎麼辦？", "濾水器需要電力或耗材嗎？", "若情境改成災後一個月，答案會變嗎？"],
        "takeaway": "勝負不在卡牌名稱，而在比較標準是否貼合情境、能否涵蓋重要後果。",
    },
    {
        "title": "6｜推理解密",
        "tag": "問題設計・分類排除・資訊效率",
        "goal": ["提出能有效縮小範圍的問題", "依答案更新推理而非盲猜", "區分特徵、類別與偶然資訊"],
        "time": "10–20 分鐘；適合全班合作或小隊競速。",
        "prepare": ["選定詞庫與候選範圍。", "老師秘密輸入答案編號，確認投影不會洩漏。", "事先約定只能回答：是、不是、不一定、接近了。"],
        "steps": ["先觀察候選詞條可以如何分類。", "前五題只能提問，不可直接猜答案。", "每次回答後，學生說明排除了哪些候選。", "範圍縮小後才允許猜測。", "公布答案並比較哪一題帶來最多資訊。"],
        "coach": ["鼓勵一題切掉大量候選，而不是只確認一個名稱。", "遇到『不一定』時，請學生修正問題的條件。", "不要只看猜中速度，也要回顧問題品質。"],
        "setup": "詞庫：動物卡｜秘密答案：企鵝",
        "answer": "第一題問『牠主要生活在陸地嗎？』可能得到不一定，資訊有限；更好的問題是『牠是鳥類嗎？』，能一次排除多數候選。接著問『牠能飛嗎？』與『牠通常生活在寒冷地區嗎？』，就能快速接近企鵝。",
        "questions": ["哪一題排除的候選最多？", "『牠可愛嗎』為什麼不是好問題？", "如果答案是『不一定』，要怎麼改寫？"],
        "takeaway": "推理能力不只是猜對，而是用少量、可判斷的問題取得最多資訊。",
    },
    {
        "title": "7｜辯論黑板",
        "tag": "正式辯論・角色分工・時間管理",
        "goal": ["讓學生辨認正反方、辯位與發言順序", "用清楚辯題進行完整攻防", "練習時間感與團隊交接"],
        "time": "20–50 分鐘；依辯士人數與賽制調整。",
        "prepare": ["輸入辯題與正反方隊名。", "設定每方 1–5 位辯士，視需要填姓名並標記結辯。", "選擇共用計時或雙方碼錶，先試一次提示鈴。"],
        "steps": ["老師先確認辯題中的主體、行動與比較方向。", "雙方分配立論、攻辯、資料與結辯責任。", "依黑板座位與計時器進行發言。", "每段結束由下一位辯士先回應前段衝突，再推進己方論點。", "賽後依主張、理由、證據、回應四項回饋。"],
        "coach": ["黑板是流程工具，不是自動產生論點的遊戲。", "發言前要求辯士說清楚本段任務，避免重複隊友內容。", "結辯只能整理已出現的衝突，不應突然提出全新論點。"],
        "setup": "辯題：國小應取消回家作業｜正方 3 人、反方 3 人",
        "answer": "正一主張取消作業能降低家庭資源差距；反一指出課後練習仍有必要。正二不否認練習價值，而是區分『在校完成的指導練習』與『依賴家長的回家作業』，把爭點收束成：學習成效是否必須透過回家作業達成。",
        "questions": ["雙方真正不同意的是作業本身，還是完成作業的地點？", "哪一方提出了可以比較的成效標準？", "哪個回應有直接處理對方理由？"],
        "takeaway": "正式辯論的核心是持續處理同一個衝突；黑板與計時器幫助學生看見角色、順序與責任。",
    },
    {
        "title": "8｜卡片字典",
        "tag": "自由組合・課程客製・學生造規則",
        "goal": ["讓老師快速挑選跨牌庫素材", "依教學目標設計臨時活動", "讓學生練習制定公平、可執行的規則"],
        "time": "5–20 分鐘；適合作為備課工具或課堂即興活動。",
        "prepare": ["先說清楚今天要練比較、故事、推理還是辯護。", "啟用需要的卡池，再直接挑選卡片。", "常用組合可儲存成預設，下次直接載入。"],
        "steps": ["老師或學生選出 3–6 張跨類別卡。", "用一句話說明這些卡要完成什麼任務。", "訂定回合、時間、成功條件與不能做的事。", "試玩一輪，記錄太容易、太難或有爭議的地方。", "修改規則後再玩一次。"],
        "coach": ["先有學習目標，再選卡，不要只因卡片有趣而堆在一起。", "學生造規則時要追問：怎樣算贏？每個人機會相同嗎？", "第一版只需能玩，試玩後再增加複雜規則。"],
        "setup": "選卡：無人島、醫生、垃圾袋、信任｜自訂活動：60 秒救援提案",
        "answer": "規則是每組必須使用四張卡提出救援計畫，並說明如何建立島上居民對醫生的信任。成功條件不是講得最誇張，而是四張卡都有功能、行動有順序，而且能回答一次質疑。",
        "questions": ["如果某張卡完全沒用到，是否扣分？", "怎麼避免口才最好的人永遠獲勝？", "要增加什麼限制才能讓第二輪更有挑戰？"],
        "takeaway": "卡片字典不是第九種固定玩法，而是讓老師把教學目標轉成可試玩的活動原型。",
    },
]


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size in ((1, 18), (2, 14), (3, 11.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY if level == 1 else BLUE
        style.paragraph_format.keep_with_next = True


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    set_run(p.add_run("辯語視界 DebateVision｜八項活動教練手冊"), size=8.5, color=MUTED)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    add_footer(section)

    # Editorial-cover pattern, adapted for a practical teacher handbook.
    add_text(doc, "辯語視界 DebateVision", size=11, bold=True, color=GOLD, after=64, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "八項活動教練手冊", size=29, bold=True, color=NAVY, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "給第一次帶領思辨活動的老師", size=15, color=BLUE, after=22, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "玩法 × 示範 × 追問 × 收束", size=11.5, bold=True, color=GOLD, after=70, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_labeled(doc, "本手冊的目的", "讓老師不只會操作網站，也能把學生的有趣回答推進成有理由、有步驟、能回應質疑的思辨表達。", LIGHT_BLUE)
    add_text(doc, "版本：2026 年 8 月｜適用於桌機與手機版現行八項活動", size=9.5, color=MUTED, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

    heading(doc, "如何使用這份手冊", 1)
    add_text(doc, "每一項活動都以同一套教練節奏設計。第一次帶課時，不必逐字照念；先掌握活動目標、回合順序與三種追問即可。", size=10.8, color=NAVY)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.05, 1.65, 1.8, 2.0]
    for i, (label, width) in enumerate(zip(["階段", "老師要做什麼", "學生要做什麼", "完成的判準"], widths)):
        table.columns[i].width = Inches(width)
        set_cell_text(table.rows[0].cells[i], label, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.rows[0].cells[i], "27698B")
    rows = [
        ("設定", "選玩法、卡池與時間", "理解限制與任務", "所有人知道這局要解決什麼"),
        ("提出", "限制準備時間", "給出主張與理由", "答案包含具體行動或判準"),
        ("追問", "挑一個核心漏洞", "補充、修正或反駁", "有直接處理問題"),
        ("收束", "比較方法而非只判輸贏", "說出學到的策略", "能指出答案為何變好"),
    ]
    for row_i, values in enumerate(rows):
        cells = table.add_row().cells
        for col_i, value in enumerate(values):
            set_cell_text(cells[col_i], value, bold=col_i == 0, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER if col_i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
            if row_i % 2 == 0:
                shade(cells[col_i], "F7FAFB")
    set_table_geometry(table, [1512, 2376, 2592, 2880])

    heading(doc, "教練的三種萬用追問", 2)
    add_bullets(doc, [
        "具體化：你在哪一個步驟使用它？誰先做什麼？",
        "找限制：這個方案最可能在哪裡失敗？需要什麼條件才成立？",
        "做比較：和另一個選項相比，為什麼這個標準更適合現在的情境？",
    ])
    add_labeled(doc, "安全提醒", "不需要羞辱答錯的人。學生若回答太誇張，教練的工作是幫他補上條件、順序與限制，而不是立刻否定創意。", LIGHT_GOLD)

    heading(doc, "八項活動快速選擇", 1)
    quick = doc.add_table(rows=1, cols=3)
    quick.alignment = WD_TABLE_ALIGNMENT.CENTER
    quick.autofit = False
    for i, (label, width) in enumerate(zip(["想練的能力", "建議活動", "最短可用時間"], [2.7, 2.6, 1.2])):
        quick.columns[i].width = Inches(width)
        set_cell_text(quick.rows[0].cells[i], label, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(quick.rows[0].cells[i], "1A374D")
    quick_rows = [
        ("限制下找用途、團隊策略", "異境求生", "10 分鐘"),
        ("把創意變成現實方案", "現實召喚", "12 分鐘"),
        ("需求、客群與說服", "銷售密令", "10 分鐘"),
        ("比喻、定義與反例", "隱喻羅盤", "8 分鐘"),
        ("比較標準與攻防", "誰更重要", "8 分鐘"),
        ("問題設計與排除法", "推理解密", "10 分鐘"),
        ("正式辯論與時間管理", "辯論黑板", "20 分鐘"),
        ("臨時選卡、自己造規則", "卡片字典", "5 分鐘"),
    ]
    for r_i, values in enumerate(quick_rows):
        cells = quick.add_row().cells
        for c_i, value in enumerate(values):
            set_cell_text(cells[c_i], value, bold=c_i == 1, align=WD_ALIGN_PARAGRAPH.CENTER if c_i > 0 else WD_ALIGN_PARAGRAPH.LEFT)
            if r_i % 2 == 0:
                shade(cells[c_i], "F3F7F9")
    set_table_geometry(quick, [3888, 3744, 1728])

    for activity in ACTIVITIES:
        heading(doc, activity["title"], 1)
        add_text(doc, activity["tag"], size=11, bold=True, color=GOLD, after=8)
        heading(doc, "教學目標", 2)
        add_bullets(doc, activity["goal"])
        add_labeled(doc, "時間與分組", activity["time"], LIGHT_BLUE)
        heading(doc, "課前準備", 2)
        add_steps(doc, activity["prepare"])
        heading(doc, "帶領流程", 2)
        add_steps(doc, activity["steps"])
        heading(doc, "教練觀察重點", 2)
        add_bullets(doc, activity["coach"])
        add_example(doc, activity["setup"], activity["answer"], activity["questions"], activity["takeaway"])

    heading(doc, "課後回饋與教練自評", 1)
    add_text(doc, "每次活動結束後，只需記錄一項做得好與一項下次調整。若學生覺得『好玩但不知道學到什麼』，通常不是活動失敗，而是缺少最後兩分鐘的收束。", size=10.8)
    heading(doc, "給學生的四句回饋", 2)
    add_bullets(doc, [
        "你的主張很清楚，下一步要補的是理由。",
        "你的理由有創意，現在請說明它在哪個條件下成立。",
        "你有回答問題，但還可以更直接處理對方最強的質疑。",
        "你修改了原本的答案，這不是退讓，而是讓論述更準確。",
    ])
    heading(doc, "教練自評清單", 2)
    add_bullets(doc, [
        "學生是否知道本局練習的能力，而不只是抽到什麼卡？",
        "準備時間是否足夠，但沒有長到讓少數人包辦？",
        "追問是否集中在一個核心問題？",
        "是否讓學生有機會修改答案？",
        "結束時是否說明哪一種思考方法值得帶到下一局？",
    ])
    add_labeled(doc, "建議", "第一次交給新教練時，可先請他只帶『誰更重要』或『異境求生』一項；熟悉提出—追問—修正—收束後，再擴充其他玩法。", LIGHT_GOLD)

    doc.core_properties.title = "辯語視界八項活動教練手冊"
    doc.core_properties.subject = "思辨教育活動玩法、示例與教師追問"
    doc.core_properties.author = "辯語視界 DebateVision"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
